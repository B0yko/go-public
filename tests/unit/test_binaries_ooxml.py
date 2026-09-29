"""Dev-only: `bench/binaries.py`'s hand-written OOXML packages load in the real
`python-docx`/`openpyxl` libraries (a dev-only test). Neither library is a runtime
dependency — `fixture`, `demo` and `bench` build these bytes with `zipfile` alone
(see that module's docstring); this test alone justifies the dev-only dependency.
"""

from __future__ import annotations

import io

import docx
import openpyxl

from go_public.bench import binaries


def test_minimal_docx_loads_in_python_docx() -> None:
    content = binaries.minimal_docx(creator="Casey Morgan")
    document = docx.Document(io.BytesIO(content))
    assert document.core_properties.author == "Casey Morgan"
    assert any("Hello." in p.text for p in document.paragraphs)


def test_minimal_docx_with_comment_and_tracked_change_loads() -> None:
    content = binaries.minimal_docx(
        tracked_change_author="Robin Taylor", comment_author="Avery Quinn"
    )
    document = docx.Document(io.BytesIO(content))
    assert document.paragraphs  # the package parses without error


def test_minimal_xlsx_loads_in_openpyxl() -> None:
    content = binaries.minimal_xlsx(company="Silverline Fictional Ltd")
    workbook = openpyxl.load_workbook(io.BytesIO(content))
    sheet = workbook["Sheet1"]
    assert sheet["A1"].value == "hello"
